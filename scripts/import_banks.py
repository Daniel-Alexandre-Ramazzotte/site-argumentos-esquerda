"""Preserva documentos e anexos oficiais de rentabilidade bancária."""
import sys
import io
import json
from import_mobility import download, ROOT
from concurrent.futures import ThreadPoolExecutor
def one(date):
    import pymupdf
    url=f'https://www.bcb.gov.br/content/publicacoes/ref/{date}/RELESTAB{date}-refPub.pdf'
    try:
        raw=download(url,f'bank-ref-{date}.pdf')
        doc=pymupdf.open(stream=raw,filetype='pdf')
        for i,p in enumerate(doc):
            text=p.get_text()
            if 'lucro líquido' in text.lower():
                print('DOC',date,'PAG',i+1,text[:2200],flush=True)
                print('LINKS',[l.get('uri') for l in p.get_links() if l.get('uri')],flush=True)
    except Exception as exc: print(date,str(exc),flush=True)
sys.stdout.reconfigure(encoding='utf-8')
if '--normalize' in sys.argv:
    import openpyxl
    from datetime import datetime
    profit=[]
    wb=openpyxl.load_workbook(ROOT/'data/raw/bank-ref-202404.xlsx',data_only=True,read_only=True)
    # Apesar do cabeçalho R$ bilhões, o anexo 202404 armazena valores em reais.
    for row in wb['Graf 1.2.3.1'].values:
        if isinstance(row[0],datetime) and row[0].month==12:
            assert row[1]>1e9
            profit.append({'year':row[0].year,'value':row[1]/1e9})
    sfn=[]
    wb=openpyxl.load_workbook(ROOT/'data/raw/bank-ref-202605.xlsx',data_only=True,read_only=True)
    for row in wb['Graf 1.2.3.1'].values:
        if isinstance(row[0],str) and len(row[0])==7 and row[0].endswith('-12') and row[0][:4].isdigit() and int(row[0][:4])<=2025:
            number=lambda v: float(v.strip().replace('.','').replace(',','.')) if isinstance(v,str) else float(v)
            sfn.append({'year':int(row[0][:4]),'profit':number(row[1])/1e9,'roe':number(row[3]),'privateS1':number(row[5])/1e9,'public':number(row[9])/1e9})
    assert len(profit)>1 and len(sfn)>1
    (ROOT/'data/bank-profit.json').write_text(json.dumps({'profit':profit,'sfn':sfn},indent=2),encoding='utf-8')
    print('Bancos',profit,'SFN',sfn,flush=True)
    sys.exit()
if '--sheets' in sys.argv:
    import openpyxl
    for date in ('202404','202605'):
        url=f'https://www.bcb.gov.br/content/publicacoes/ref/{date}/RELESTAB{date}-refAnexoEstatistico.xlsx'
        raw=download(url,f'bank-ref-{date}.xlsx')
        wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
        for sheet in wb:
            rows=list(sheet.values)
            if any('lucro' in str(r).lower() for r in rows) and sheet.title.startswith('Graf 1.2.3'):
                print(date,sheet.title,rows[:16],flush=True)
    sys.exit()
list(ThreadPoolExecutor(max_workers=3).map(one,('202304','202504','202605')))
