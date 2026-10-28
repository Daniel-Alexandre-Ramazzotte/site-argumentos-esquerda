"""Séries oficiais: crédito às famílias e documentos empresariais. Preserva originais."""
import json
import sys
import io
import re
import zipfile
import warnings
from concurrent.futures import ThreadPoolExecutor
from import_mobility import download, ROOT

def credit():
    def fetch(code):
        url = f'https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados?formato=json&dataInicial=01/01/1985&dataFinal=30/09/2026'
        rows = json.loads(download(url, f'credit-bcb-{code}-2026-09.json'))
        monthly = sorted([{'month': r['data'][6:] + '-' + r['data'][3:5], 'value': float(r['valor'])} for r in rows], key=lambda p: p['month'])
        assert len(monthly) == len(set(p['month'] for p in monthly))
        points = [{'year': int(p['month'][:4]), 'value': p['value']} for p in monthly if p['month'].endswith('-12') and p['month'] < '2026']
        assert len(points) > 1
        print(code, len(monthly), monthly[0], monthly[-1], flush=True)
        return str(code), {'source': url, 'monthly': monthly, 'points': points}
    data = dict(ThreadPoolExecutor(max_workers=4).map(fetch, (4189, 20716, 29037, 29034, 29038)))
    (ROOT/'data'/'household-credit.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

def documents():
    urls = {
        'health-upa.zip': 'https://demas-dados-abertos.s3.amazonaws.com/json/upacrm.json.zip',
        'health-swagger.json': 'https://apidadosabertos.saude.gov.br/static/swagger.json',
        'rfb-2024.xlsx': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-2024/tabela-08-forma-de-tributacao/@@download/file',
        'rfb-2021-2022.xlsx': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-consolidados-2021-e-2022/tabelas-resumo/@@download/file',
        'rfb-2023-page.html': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-2023',
        'rfb-2019-page.html': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-consolidados-2019-e-2020',
        'rfb-2016-page.html': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-consolidados-2016-a-2018',
        'mapa-2020.pdf': 'https://www.gov.br/empresas-e-negocios/pt-br/mapa-de-empresas/boletins/mapa-de-empresas-boletim-do-3o-quadrimestre-de-2020.pdf/@@download/file',
        'mapa-2021.pdf': 'https://www.gov.br/empresas-e-negocios/pt-br/mapa-de-empresas/boletins/mapa-de-empresas-boletim-do-3o-quadrimestre-de-2021.pdf/@@download/file',
        'mapa-2022.pdf': 'https://www.gov.br/empresas-e-negocios/pt-br/mapa-de-empresas/boletins/mapa-de-empresas-boletim-3o-quadrimestre-2022.pdf/@@download/file',
        'mapa-2023.pdf': 'https://www.gov.br/empresas-e-negocios/pt-br/mapa-de-empresas/boletins/mapa-de-empresas-boletim-3o-quadrimestre-2023.pdf/@@download/file',
        'mapa-2024.pdf': 'https://www.gov.br/empresas-e-negocios/pt-br/mapa-de-empresas/boletins/boletim-do-mapa-de-empresas-3o-quad-2024.pdf/@@download/file',
        'rfb-2023.xlsx': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-2023/tab-08-ft.xlsx/@@download/file',
        'rfb-2019-2020.xlsx': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-consolidados-2019-e-2020/tabelas-resumo.xlsx/@@download/file',
        'rfb-2016-2018.xlsx': 'https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/dados-setoriais-consolidados-2016-a-2018/tabela-4-dados-consolidados-2016-a-2018-brasil-e-estados/@@download/file',
    }
    def one(item):
        name, url = item
        try:
            raw = download(url, name)
            print(name, len(raw), raw[:30], flush=True)
            if name.endswith('.html'):
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(raw, 'html.parser')
                print(name, [(a.text.strip(),a.get('href')) for a in soup.select('#content-core a')], flush=True)
            if name.endswith('.pdf'):
                import pymupdf
                doc = pymupdf.open(stream=raw, filetype='pdf')
                text = '\n'.join(p.get_text() for p in doc)
                (ROOT/'data'/'raw'/name.replace('.pdf','.txt')).write_text(text, encoding='utf-8')
                print(name, re.findall(r'.{0,100}empresas ativas.{0,100}', text, re.I)[:4], flush=True)
            if name.endswith('.xlsx'):
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
                for sheet in wb:
                    print(name, sheet.title, list(sheet.values)[:7], flush=True)
            if name.endswith('.zip'):
                with zipfile.ZipFile(io.BytesIO(raw)) as z:
                    for f in z.namelist():
                        rows = json.loads(z.read(f))
                        print(f, type(rows), str(rows)[:1800], flush=True)
        except Exception as exc:
            print(name, type(exc).__name__, str(exc), flush=True)
    list(ThreadPoolExecutor(max_workers=4).map(one, urls.items()))

def health():
    with zipfile.ZipFile(ROOT/'data/raw/health-upa.zip') as z:
        rows = json.loads(z.read(z.namelist()[0]))
    # Os totais de Brasil/UF/região se repetem em cada município; nunca somá-los.
    geos = {}
    for r in rows:
        month = str(r['co_anomes'])
        if not month.endswith('12'): continue
        for ident, name, level, value in [
            ('br','Brasil','br',r['vl_indicador_calculado_br']),
            (f"state-{r['co_uf']}",r['no_uf'],'state',r['vl_indicador_calculado_uf']),
            (f"region-{r['co_regiao_brasil']}",r['no_regiao_brasil'],'region',r['vl_indicador_calculado_reg']),
        ]:
            geo = geos.setdefault(ident, {'id':ident,'name':name,'level':level,'values':{}})
            if month in geo['values']: assert geo['values'][month] == value
            geo['values'][month] = value
    output = []
    for geo in geos.values():
        values = geo.pop('values')
        geo['points'] = [{'year':int(m[:4]),'value':v} for m,v in sorted(values.items())]
        output.append(geo)
    (ROOT/'data/health-infrastructure.json').write_text(json.dumps({'upa':output},ensure_ascii=False,indent=2),encoding='utf-8')
    print('UPA',len(output),'geografias',output[0]['points'],flush=True)

def business():
    import openpyxl
    warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    regimes = {k: {'count': [], 'revenue': []} for k in ('LUCRO REAL', 'LUCRO PRESUMIDO')}
    old = openpyxl.load_workbook(ROOT/'data/raw/rfb-2016-2018.xlsx', data_only=True, read_only=True)
    national = list(old['BR'].values)
    for name in regimes:
        revenue = next(r for r in national[:20] if r[1] == name)
        count = next(r for r in national[43:62] if r[1] == name)
        for year, col in [(2016,2),(2017,4),(2018,6)]:
            regimes[name]['revenue'].append({'year':year,'value':revenue[col]*1e6})
            regimes[name]['count'].append({'year':year,'value':count[col]})
    for filename, years in [('rfb-2019-2020.xlsx',(2019,2020)),('rfb-2021-2022.xlsx',(2021,2022))]:
        sheet = list(openpyxl.load_workbook(ROOT/'data/raw'/filename,data_only=True,read_only=True)['BR'].values)
        for name in regimes:
            revenue = next(r for r in sheet[6:12] if r[1] == name)
            count = next(r for r in sheet[6:12] if r[7] == name)
            for year,col in zip(years,(2,4)):
                regimes[name]['revenue'].append({'year':year,'value':revenue[col]*1e6})
                regimes[name]['count'].append({'year':year,'value':count[col+6]})
    for year in (2023,2024):
        wb = openpyxl.load_workbook(ROOT/f'data/raw/rfb-{year}.xlsx',data_only=True,read_only=True)
        rows = list(wb.active.values)
        for name in regimes:
            row = next(r for r in rows if r[0] == name)
            regimes[name]['count'].append({'year':year,'value':row[1]})
            regimes[name]['revenue'].append({'year':year,'value':row[2]})
    stocks=[]
    for year in range(2020,2025):
        text=(ROOT/f'data/raw/mapa-{year}.txt').read_text(encoding='utf-8')
        # O resumo e a apresentação repetem o total. Confirmar o valor nas duas seções.
        matches=re.findall(r'([\d.]+)\s+empresas ativas',text[:15000])
        assert matches and len(set(matches))==1, (year,matches)
        value=int(matches[0].replace('.',''))
        assert value>10_000_000
        stocks.append({'year':year,'value':value})
    (ROOT/'data/business.json').write_text(json.dumps({'regimes':regimes,'activeSnapshots':stocks},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'regimes':regimes,'activeSnapshots':stocks},ensure_ascii=False),flush=True)

if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    if '--documents' in sys.argv: documents()
    elif '--health' in sys.argv: health()
    elif '--business' in sys.argv: business()
    else: credit()
