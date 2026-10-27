"""Importa fontes oficiais para câmbio, combustíveis e transporte; mantém os originais."""
import gzip
import json
import sys
import io
import re
import unicodedata
from datetime import datetime
from collections import defaultdict
from pathlib import Path
from urllib.request import Request, urlopen
from concurrent.futures import ThreadPoolExecutor

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data' / 'raw'
RAW.mkdir(parents=True, exist_ok=True)

def download(url, filename):
    path = RAW / filename
    if path.exists():
        return path.read_bytes()
    data = urlopen(Request(url, headers={'User-Agent': 'Mozilla/5.0'}), timeout=90).read()
    if data.startswith(b'\x1f\x8b'):
        data = gzip.decompress(data)
    path.write_bytes(data)
    return data

def inspect():
    import io
    import openpyxl
    base = 'https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/precos-revenda-e-de-distribuicao-combustiveis/shlp/'
    for name, suffix in [('old-br', '2001-2012/mensal-brasil-2001-a-2012.xlsx'), ('new-br','mensal/mensal-brasil-desde-jan2013.xlsx'), ('new-state','mensal/mensal-estados-desde-jan2013.xlsx')]:
        data = download(base + suffix, f'anp-{name}.xlsx')
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        print(name, wb.sheetnames)
        for sheet in wb:
            print(sheet.title, list(sheet.values)[:14])
    for code in (3698, 3695):
        url = f'https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados?formato=json&dataInicial=01/07/1994&dataFinal=31/12/2025'
        rows = json.loads(download(url, f'bcb-{code}.json'))
        print('BCB',code, len(rows), rows[:2],rows[-2:])
    for code in ('IS.RRS.TOTL.KM', 'IS.AIR.PSGR', 'IS.SHP.GOOD.TU'):
        url = f'https://api.worldbank.org/v2/country/BRA/indicator/{code}?format=json&per_page=1000'
        data = json.loads(download(url,f'worldbank-{code}.json'))
        rows = [r for r in data[1] if r['value'] is not None]
        print('WB', code, len(rows), [(r['date'],r['value']) for r in rows[:3]],[(r['date'],r['value']) for r in rows[-3:]])

def key(value):
    return re.sub(r'[^A-Z0-9]', '', unicodedata.normalize('NFKD', str(value)).encode('ascii', 'ignore').decode().upper())

def exchange():
    result = {}
    for code in (3698,3695):
        url=f'https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados?formato=json&dataInicial=01/07/1994&dataFinal=30/09/2026'
        rows=json.loads(download(url,f'bcb-{code}-through-2026-09.json'))
        monthly = [{'month':r['data'][-4:]+'-'+r['data'][3:5], 'value':float(r['valor'])} for r in rows]
        monthly.sort(key=lambda r:r['month'])
        if len(monthly) != len(set(r['month'] for r in monthly)):
            raise ValueError('Mês duplicado na série cambial')
        points=[{'year':int(r['month'][:4]),'value':r['value']} for r in monthly if r['month'][5:]=='12' and int(r['month'][:4])<=2025]
        result[str(code)]={'points':points,'monthly':monthly,'source':url}
        print('BCB',code,len(monthly),monthly[0],monthly[-1],flush=True)
    return result

def inflation():
    url='https://apisidra.ibge.gov.br/values/t/1737/n1/1/v/2266/p/all'
    rows=json.loads(download(url,'sidra-ipca-number-index-through-2026-09.json'))
    points=[{'month':r['D3C'][:4]+'-'+r['D3C'][4:6], 'index':float(r['V'])} for r in rows[1:] if '199407'<=r['D3C']<='202609' and r['V'] not in ('...','-','..')]
    points.sort(key=lambda r:r['month'])
    base=next(r['index'] for r in points if r['month']=='2025-12')
    factors=[{**r,'factor':base/r['index']} for r in points]
    if any(r['index']<=0 for r in factors) or len(factors)!=387:
        raise ValueError(f'Índice mensal incompleto ou inválido: {len(factors)}')
    result={'baseMonth':'2025-12','source':url,'points':factors}
    (ROOT/'data/price-index.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print('IPCA número-índice',len(factors),factors[0],factors[-1],flush=True)

def build():
    import openpyxl
    page = 'https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/precos-revenda-e-de-distribuicao-combustiveis/serie-historica-do-levantamento-de-precos'
    html = download(page, 'anp-monthly-index.html').decode('utf-8')
    urls = list(dict.fromkeys(re.findall(r'href="([^"]+\.xlsx)"', html)))
    urls = [url for url in urls if 'mensal-' in url and any(t in url for t in ('brasil', 'regioes', 'estados'))]
    if len(urls) != 6:
        raise ValueError(f'Esperadas seis planilhas mensais, encontradas {len(urls)}: {urls}')
    geography = json.loads((ROOT/'data/geography.json').read_text(encoding='utf-8'))['unemployment']
    territories = {key(g['name']): {k:g[k] for k in ('id','name','level')} for g in geography}
    territories['BRASIL'] = {'id':'br','name':'Brasil','level':'br'}
    observations = defaultdict(dict)
    sources = []
    for url in urls:
        payload = download(url, 'fuel-' + url.rsplit('/', 1)[-1])
        sheet = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True).active
        headers = None
        count = 0
        for row in sheet.values:
            if row and key(row[0]) == 'MES':
                headers = {key(v):i for i,v in enumerate(row) if v is not None}
                continue
            if not headers or not isinstance(row[0], datetime):
                continue
            if row[0].month != 12 or row[0].year > 2025:
                continue
            product = key(row[headers['PRODUTO']])
            place = row[headers['ESTADO']] if 'ESTADO' in headers else row[headers['REGIAO']] if 'REGIAO' in headers else 'Brasil'
            territory = territories.get(key(place))
            if not territory:
                raise ValueError(f'Território não reconhecido: {place}')
            price = row[headers['PRECOMEDIOREVENDA']]
            if not isinstance(price,(int,float)) or price <= 0:
                continue
            group = (product, territory['id'])
            year = row[0].year
            if year in observations[group]:
                raise ValueError(f'Observação duplicada: {group}, {year}')
            observations[group][year] = round(float(price),4)
            count += 1
        sources.append(url)
        print(url.rsplit('/',1)[-1],count,flush=True)
    fuels = {}
    for product in sorted(set(p for p,t in observations)):
        geos = []
        for territory in territories.values():
            rows = observations.get((product,territory['id']),{})
            if rows:
                geos.append({**territory,'points':[{'year':y,'value':v} for y,v in sorted(rows.items())]})
        fuels[product] = geos
    currency = exchange()
    # Volumes operacionais são mantidos separados de extensão da infraestrutura.
    transport = {}
    for code in ('IS.AIR.PSGR','IS.SHP.GOOD.TU'):
        url=f'https://api.worldbank.org/v2/country/BRA/indicator/{code}?format=json&per_page=1000'
        rows=json.loads(download(url,f'worldbank-{code}.json'))[1]
        points=sorted([{'year':int(r['date']),'value':float(r['value'])} for r in rows if r['value'] is not None and int(r['date'])<=2025],key=lambda r:r['year'])
        transport[code]={'points':points,'source':url}
    result={'fuels':fuels,'fuelSources':sources,'currency':currency,'transport':transport}
    (ROOT/'data/mobility.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    for product,geos in fuels.items():
        br=next(g for g in geos if g['id']=='br')
        print(product,len(geos),br['points'][0],br['points'][-1],flush=True)

def documents():
    import fitz
    url='https://ontl.infrasa.gov.br/wp-content/uploads/2026/01/Diagnostico_logistico-2023-database_2010_2022-ajustado-1.pdf'
    payload=download(url,'ontl-diagnostic-2010-2022.pdf')
    doc=fitz.open(stream=payload,filetype='pdf')
    for page in (18,36,37,54,58):
        print('PAGE',page+1,doc[page].get_text(),flush=True)
        dest=ROOT/'artifacts'/f'transport-source-{page+1}.png'
        dest.parent.mkdir(exist_ok=True)
        doc[page].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(dest)

if __name__ == '__main__':
    if '--inspect' in sys.argv:
        inspect()
    elif '--documents' in sys.argv:
        documents()
    elif '--currency-only' in sys.argv:
        path = ROOT/'data/mobility.json'
        result=json.loads(path.read_text(encoding='utf-8'))
        result['currency']=exchange()
        path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    elif '--inflation-only' in sys.argv:
        inflation()
    else:
        build()
